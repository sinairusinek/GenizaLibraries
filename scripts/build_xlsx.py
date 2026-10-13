"""Build allony_data.xlsx: the data behind the page, one sheet per table."""
import json, csv
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
E = json.load(open("data/allony_index.json")); C = json.load(open("data/fgp_categories.json")); L = json.load(open("data/allony_lists.json"))
CAT = {c["code"]: c for c in C}; LBY = {l["list"]: l for l in L}
def lvl(code, n):
    p = code.split("."); return ".".join(p[:n]) if len(p) >= n else ""
wb = Workbook()
def sheet(title, header, widths, rtl_cols=()):
    ws = wb.create_sheet(title); ws.append(header); ws.freeze_panes = "A2"
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]: c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}1"
    return ws
ws = wb.active; ws.title = "About"
for row in [["Allony index, classed by FGP field"], [""],
            ["Source 1", "Index to N. Allony, The Jewish Library in the Middle Ages: Book Lists from the Cairo Genizah, ed. M. Frenkel and H. Ben-Shammai (Jerusalem: Ben-Zvi Institute, 2006), pp. 457-508."],
            ["Source 2", "Friedberg Genizah Project, list of fields with statistics, as printed 9 Oct 2026."],
            ["Source 3", "Princeton Geniza Project, public document metadata (dates, descriptions), matched by shelfmark."],
            ["Caution", "The FGP class of each entry is a suggestion made from general knowledge of the literature. It was not checked against the book lists."],
            ["Caution", "Dates and places marked 'inferred' come from the person named in the list, not from the book."],
            ["Sheets", "Index entries · Mentions by list · Book lists · FGP categories · Cross-references"],
            ["Mention", "One line reference of an index entry to a book list. A line range (e.g. 1-25) counts as one mention."]]:
    ws.append(row)
ws["A1"].font = Font(bold=True, size=14); ws.column_dimensions["A"].width = 12; ws.column_dimensions["B"].width = 130
# ---- index entries
ws = sheet("Index entries", ["No.", "Index entry", "Gloss", "Attributed to", "Type", "FGP code", "Field", "Sub-field", "Sub-sub-field", "Confidence", "Lists", "Mentions",
                             "References (list: lines)", "See", "See also", "Note", "Highlighted in the Word file", "Index line as given"],
           [7, 38, 28, 24, 12, 9, 30, 28, 28, 10, 7, 9, 50, 30, 30, 40, 12, 70])
for e in E:
    k = e["classification"]; code = k["fgp_code"]
    refs = "; ".join(f'{r["list"]}: {", ".join(r["lines"])}' + (f' ({r["note"]})' if r["note"] else "") for r in e["refs"])
    ws.append([e["id"], e["headword"], e["gloss"], e["attributed_to"], k["type"], code,
               CAT[lvl(code, 1)]["name"] if code else "", CAT[lvl(code, 2)]["name"] if lvl(code, 2) else "", CAT[lvl(code, 3)]["name"] if lvl(code, 3) else "",
               k["confidence"], e["n_lists"], e["n_refs"], refs, " | ".join(e["see"]), " | ".join(e["see_also"]),
               (k["note"] + (" · " + k["rule"] if k["rule"] else "")).strip(" ·"), e["highlight"], e["raw"]])
    r = ws.max_row
    for n, col in ((1, 7), (2, 8), (3, 9)):
        kk = lvl(code, n)
        if kk:
            c = ws.cell(r, col); c.fill = PatternFill("solid", fgColor=CAT[kk]["color"][1:].upper()); c.font = Font(color=CAT[kk]["ink"][1:].upper())
# ---- mentions by list
ws = sheet("Mentions by list", ["List", "Shelfmark", "Type of list", "Entry no.", "Index entry", "Entry type", "FGP code", "Field", "FGP class (full)", "Confidence", "Line(s)", "Note in the index"],
           [6, 26, 30, 9, 38, 12, 9, 30, 55, 10, 12, 24])
for e in E:
    k = e["classification"]; code = k["fgp_code"]
    for r in e["refs"]:
        b = LBY.get(r["list"])
        for ln in r["lines"]:
            ws.append([r["list"], b["shelfmark"] if b else "", f'{b["type_code"]}. {b["type_en"]}' if b else "", e["id"], e["headword"], k["type"], code,
                       CAT[lvl(code, 1)]["name"] if code else "", k["fgp_path"], k["confidence"], ln, r["note"]])
# ---- book lists
ws = sheet("Book lists", ["List", "Shelfmark", "Type code", "Type of list", "Type (Hebrew)", "Sub-group in the book", "Description in the book's contents", "Page",
                          "Date: book contents", "Date: Princeton Geniza Project", "Date: inferred", "Place: book contents", "Place: Princeton Geniza Project", "Place: inferred",
                          "Index entries", "Mentions", "Princeton record", "Princeton shelfmark", "Princeton description"],
           [6, 30, 6, 32, 30, 40, 50, 7, 20, 50, 40, 26, 30, 14, 9, 9, 44, 30, 100])
for l in L:
    ws.append([l["list"], l["shelfmark"], l["type_code"], l["type_en"], l["type_he"], l["subgroup_he"], l["description_he"], l["page"], l["date_stated"], l["pgp_date"], l["date_inferred"],
               l["place_stated"], l["pgp_place"], l["place_inferred"], l["index_entries"], l["index_references"], l["pgp_url"], l["pgp_shelfmark"], l["pgp_description"]])
# ---- FGP categories
ws = sheet("FGP categories", ["Code", "Level", "Parent", "Name", "Full path", "Shelf numbers", "Units", "Residue row", "Colour"], [9, 6, 8, 44, 70, 13, 10, 11, 10])
for c in C:
    ws.append([c["code"], c["level"], c["parent"], c["name"], c["path"], c["shelfmarks"], c["units"], "yes" if c["residual"] else "", c["color"]])
    cell = ws.cell(ws.max_row, 9); cell.fill = PatternFill("solid", fgColor=c["color"][1:].upper()); cell.font = Font(color=c["ink"][1:].upper())
# ---- cross-references
byid = {e["id"]: e for e in E}
ws = sheet("Cross-references", ["From entry no.", "From entry", "Kind", "Target as written", "How matched", "Target entry no(s).", "Target entry"], [10, 38, 10, 38, 14, 16, 60])
for e in E:
    for x in e["xrefs"]:
        ws.append([e["id"], e["headword"], "see" if x["type"] == "see" else "see also", x["target"], x["match"], " ".join(map(str, x["ids"])), " | ".join(byid[t]["headword"] for t in x["ids"])])
for ws in wb.worksheets[1:]:
    ws.sheet_view.rightToLeft = False
wb.save("allony_data.xlsx")
import os; print("allony_data.xlsx", os.path.getsize("allony_data.xlsx") // 1024, "KB;", {w.title: w.max_row - 1 for w in wb.worksheets[1:]})
